#!/usr/bin/env python3
"""
Gera `inventario-indicadores.xlsx`, o formulário de entrada do painel.

    python3 tools/criar-inventario.py

O arquivo tem cinco abas:
  Instruções      como preencher, com contador de progresso
  1. Indicadores  um indicador por linha — vira um cartão do farol
  2. Componentes  os dados brutos necessários, e se você tem cada um
  3. Hierarquia   empresas, diretorias e áreas, para os filtros
  4. Tela a tela  de-para entre a apresentação atual e o painel

Depois de preenchido, `python3 tools/importar-inventario.py` converte a aba de
indicadores nas definições que o painel consome.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent

# ------------------------------------------------------------------ #
# Estilo
# ------------------------------------------------------------------ #
FONTE = "Arial"
TINTA = "1F2933"
AZUL = "1F4E79"

F_TITULO = Font(name=FONTE, size=14, bold=True, color=AZUL)
F_SUB = Font(name=FONTE, size=10, color="6B7A87")
F_CAB = Font(name=FONTE, size=9, bold=True, color="FFFFFF")
F_TXT = Font(name=FONTE, size=10, color=TINTA)
F_EX = Font(name=FONTE, size=10, italic=True, color="7F8C93")
F_NEG = Font(name=FONTE, size=10, bold=True, color=TINTA)
F_PREENCHER = Font(name=FONTE, size=10, color="0000FF")  # azul = você preenche

P_CAB = PatternFill("solid", fgColor=AZUL)
P_PREENCHER = PatternFill("solid", fgColor="FFF9DB")     # amarelo = campo de entrada
P_PRONTO = PatternFill("solid", fgColor="EEF2F5")        # cinza = já preenchido
P_EX = PatternFill("solid", fgColor="F7F9FA")

FIO = Side(style="thin", color="D5DBE0")
BORDA = Border(left=FIO, right=FIO, top=FIO, bottom=FIO)

ALTO = Alignment(vertical="top", wrap_text=True)
CENTRO = Alignment(vertical="center", horizontal="center", wrap_text=True)


def cabecalho(ws, colunas, linha=1):
    """Escreve a linha de cabeçalho e ajusta larguras."""
    for i, (titulo, largura) in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=titulo)
        c.font = F_CAB
        c.fill = P_CAB
        c.alignment = CENTRO
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha].height = 34
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def lista(ws, opcoes, col, primeira, ultima, titulo):
    """Aplica dropdown numa faixa de células.

    O Excel limita a lista embutida a 255 caracteres, incluindo as vírgulas.
    Passar disso faz o arquivo não abrir, então a checagem é dura de propósito.
    """
    texto = ",".join(opcoes)
    if len(texto) > 253:
        raise ValueError(f"Lista '{titulo}' tem {len(texto)} caracteres (máx. 253): {texto}")
    dv = DataValidation(
        type="list",
        formula1='"' + texto + '"',
        allow_blank=True,
        showDropDown=False,
        errorTitle="Valor fora da lista",
        error="Escolha uma das opções do menu.",
        promptTitle=titulo,
        prompt="Escolha uma das opções.",
    )
    ws.add_data_validation(dv)
    dv.add(f"{col}{primeira}:{col}{ultima}")


LINHAS_VAZIAS = 40

# ------------------------------------------------------------------ #
wb = Workbook()

# ================================================================== #
# Aba: Instruções
# ================================================================== #
ws = wb.active
ws.title = "Instruções"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 88
ws.column_dimensions["D"].width = 14

ws["B2"] = "Inventário de indicadores de RH"
ws["B2"].font = F_TITULO
ws["B3"] = "Formulário de entrada do Farol de KPIs. Preencha as abas 1 a 4; os campos em fundo amarelo são seus."
ws["B3"].font = F_SUB

blocos = [
    ("Como usar", None),
    (
        "1. Indicadores",
        "Um indicador por linha — cada linha vira um cartão do farol. Comece pelos que você "
        "realmente apresenta hoje, inclusive os que acha que deveriam sair. A linha 3 é um "
        "exemplo preenchido: apague ou sobrescreva.",
    ),
    (
        "2. Componentes",
        "O painel não guarda o indicador pronto, guarda os dados brutos que o formam. É o que "
        "permite filtrar por área sem calcular média de médias. Marque o que você tem e onde está.",
    ),
    (
        "3. Hierarquia",
        "As empresas, diretorias e áreas pelas quais você precisa filtrar. Uma linha por área "
        "(nível mais baixo), repetindo empresa e diretoria.",
    ),
    (
        "4. Tela a tela",
        "O de-para: cada slide ou aba de hoje e o que ele vira no painel. A aba já vem com um "
        "rascunho baseado no formato usual de apresentação semanal — corrija para a sua realidade.",
    ),
    ("Conceitos que mudam o preenchimento", None),
    (
        "Comportamento (aba 1)",
        "Pergunta sobre o INDICADOR: o valor cresce se o período for maior? Rotatividade, custo "
        "por colaborador e horas de treinamento crescem — são reexpressos em base mensal, o que "
        "torna semana, mês e ano comparáveis entre si e com a meta. Absenteísmo, horas extras em "
        "%, headcount e e-NPS não crescem: já são taxa ou foto.",
    ),
    (
        "Fluxo x estoque (aba 2)",
        "Pergunta sobre o COMPONENTE, e é outra coisa. Fluxo acontece ao longo do período e soma: "
        "admissões, desligamentos, horas, custo. Estoque é foto do fim do período e não soma no "
        "tempo: headcount, vagas em aberto, posições de liderança. Cuidado com a mistura: "
        "absenteísmo é uma taxa que independe do período (aba 1), mas é formada por horas, que "
        "são fluxo e somam (aba 2).",
    ),
    (
        "Direção",
        "'Menor é melhor' para rotatividade, absenteísmo, custo. 'Maior é melhor' para "
        "treinamento, e-NPS, diversidade. 'Perto da meta' para headcount contra o orçado — "
        "ficar abaixo não é economia, é posição não preenchida.",
    ),
    (
        "Meta é por",
        "Se a meta do grupo não faz sentido ao filtrar uma empresa, marque o nível certo. "
        "Custo por colaborador é o caso clássico: a média do grupo é inútil quando o mix de "
        "cargos muda entre indústria, varejo e tecnologia.",
    ),
    (
        "Tolerância",
        "Onde o amarelo virá vermelho, em % de desvio da meta. Verde é dentro da meta; "
        "amarelo é fora da meta mas dentro da tolerância; vermelho é acima dela.",
    ),
    (
        "Periodicidade real",
        "Se o dado é trimestral (clima, e-NPS), diga. Na leitura semanal ele aparece como "
        "'não apurado' em cinza — nunca como zero, nem repetindo o último valor.",
    ),
]

r = 5
for titulo, texto in blocos:
    if texto is None:
        ws.cell(row=r, column=2, value=titulo).font = Font(name=FONTE, size=11, bold=True, color=AZUL)
        r += 1
        continue
    ws.cell(row=r, column=2, value=titulo).font = F_NEG
    ws.cell(row=r, column=2).alignment = ALTO
    c = ws.cell(row=r, column=3, value=texto)
    c.font = F_TXT
    c.alignment = ALTO
    ws.row_dimensions[r].height = 15 * (1 + len(texto) // 95)
    r += 1

r += 1
ws.cell(row=r, column=2, value="Progresso").font = Font(name=FONTE, size=11, bold=True, color=AZUL)
r += 1
progresso = [
    # colunas: C=nome, L=meta é por, O=estratégico | Componentes D=você tem
    ("Indicadores preenchidos", "=COUNTA('1. Indicadores'!C4:C200)"),
    ("Componentes respondidos", "=COUNTIF('2. Componentes'!D3:D200,\"<>\")"),
    ("Áreas mapeadas", "=COUNTA('3. Hierarquia'!D4:D200)"),
    ("Telas mapeadas", "=COUNTA('4. Tela a tela'!D3:D200)"),
    ("Indicadores sem meta definida", "=COUNTIF('1. Indicadores'!L4:L200,\"Não tem meta\")"),
    ("Indicadores marcados como estratégicos", "=COUNTIF('1. Indicadores'!O4:O200,\"Sim\")"),
]
for rotulo, formula in progresso:
    ws.cell(row=r, column=2, value=rotulo).font = F_TXT
    c = ws.cell(row=r, column=4, value=formula)
    c.font = F_NEG
    c.alignment = CENTRO
    c.fill = P_PRONTO
    c.border = BORDA
    r += 1

r += 1
ws.cell(row=r, column=2, value="Legenda de cores").font = Font(name=FONTE, size=11, bold=True, color=AZUL)
r += 1
for cor, fill, txt in [
    ("amarelo", P_PREENCHER, "campo que você preenche"),
    ("cinza", P_PRONTO, "já preenchido ou calculado — não precisa mexer"),
    ("itálico", P_EX, "linha de exemplo, para apagar"),
]:
    c = ws.cell(row=r, column=2, value=cor)
    c.fill = fill
    c.font = F_TXT
    c.border = BORDA
    ws.cell(row=r, column=3, value=txt).font = F_TXT
    r += 1

# ================================================================== #
# Aba 1: Indicadores
# ================================================================== #
ws = wb.create_sheet("1. Indicadores")
COLS = [
    ("Onde aparece hoje\n(slide, aba, relatório)", 26),
    ("Nome do indicador", 30),
    ("Grupo", 16),
    ("Como você calcula hoje\n(em palavras)", 40),
    ("Numerador\n(o que se conta em cima)", 24),
    ("Denominador\n(o que se conta embaixo)", 24),
    ("Unidade", 12),
    ("Direção", 18),
    ("Comportamento\n(o valor cresce se o período for maior?)", 26),
    ("Meta", 10),
    ("Meta é por", 14),
    ("Tolerância\n(% de desvio)", 12),
    ("Periodicidade real", 15),
    ("Estratégico?", 12),
    ("Fonte do dado\n(sistema)", 24),
    ("Observações", 34),
]
# coluna A é o número da linha
ws.cell(row=1, column=1, value="#").font = F_CAB
ws.cell(row=1, column=1).fill = P_CAB
ws.cell(row=1, column=1).alignment = CENTRO
ws.column_dimensions["A"].width = 4
for i, (t, w) in enumerate(COLS, start=2):
    c = ws.cell(row=1, column=i, value=t)
    c.font = F_CAB
    c.fill = P_CAB
    c.alignment = CENTRO
    c.border = BORDA
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 42
ws.freeze_panes = "C2"

ws.cell(row=2, column=2, value="⬇ exemplo preenchido — apague esta linha").font = F_EX

EXEMPLO = [
    "Slide 7 — Rotatividade",
    "Rotatividade voluntária",
    "Movimentação",
    "Desligamentos a pedido dividido pelo headcount médio do mês, em %",
    "Desligamentos com motivo 'pedido de demissão'",
    "Headcount médio do período",
    "%",
    "Menor é melhor",
    "Cresce com o período",
    1.45,
    "Grupo",
    25,
    "Semanal",
    "Sim",
    "Folha de pagamento",
    "Hoje só reporto o turnover total; separar o voluntário é o ganho principal",
]
for j, v in enumerate(EXEMPLO, start=2):
    c = ws.cell(row=3, column=j, value=v)
    c.font = F_EX
    c.fill = P_EX
    c.alignment = ALTO
    c.border = BORDA
ws.cell(row=3, column=1, value=1).font = F_EX
ws.row_dimensions[3].height = 56

PRIMEIRA, ULTIMA = 4, 3 + LINHAS_VAZIAS
for r in range(PRIMEIRA, ULTIMA + 1):
    ws.cell(row=r, column=1, value=r - 2).font = F_TXT
    ws.cell(row=r, column=1).alignment = CENTRO
    for j in range(2, len(COLS) + 2):
        c = ws.cell(row=r, column=j)
        c.font = F_PREENCHER
        c.fill = P_PREENCHER
        c.alignment = ALTO
        c.border = BORDA

lista(ws, ["Quadro", "Movimentação", "Presença", "Custo", "Atração", "Desenvolvimento",
           "Segurança", "Engajamento", "Diversidade", "Produtividade", "Outro"],
      "D", PRIMEIRA, ULTIMA, "Grupo")
lista(ws, ["%", "R$", "dias", "horas", "pessoas", "vagas", "pontos", "índice", "razão"],
      "H", PRIMEIRA, ULTIMA, "Unidade")
lista(ws, ["Menor é melhor", "Maior é melhor", "Perto da meta"],
      "I", PRIMEIRA, ULTIMA, "Direção")
lista(ws, ["Cresce com o período", "Independe do período"],
      "J", PRIMEIRA, ULTIMA, "Comportamento")
lista(ws, ["Grupo", "Empresa", "Diretoria", "Área", "Não tem meta"],
      "L", PRIMEIRA, ULTIMA, "Meta é por")
lista(ws, ["Semanal", "Mensal", "Trimestral", "Semestral", "Anual", "Eventual"],
      "N", PRIMEIRA, ULTIMA, "Periodicidade real")
lista(ws, ["Sim", "Não"], "O", PRIMEIRA, ULTIMA, "Estratégico?")

# ================================================================== #
# Aba 2: Componentes
# ================================================================== #
ws = wb.create_sheet("2. Componentes")
cabecalho(ws, [
    ("Componente", 22),
    ("O que é", 46),
    ("Tipo", 11),
    ("Você tem?", 12),
    ("Sistema de origem", 24),
    ("Menor granularidade\nde tempo disponível", 18),
    ("Menor quebra\norganizacional disponível", 20),
    ("Observações", 34),
])

COMPONENTES = [
    ("hc_fim", "Colaboradores ativos no último dia da semana", "Estoque"),
    ("hc_orcado", "Quadro previsto no orçamento", "Estoque"),
    ("adm", "Admissões", "Fluxo"),
    ("desl", "Desligamentos totais", "Fluxo"),
    ("desl_vol", "Desligamentos a pedido do colaborador", "Fluxo"),
    ("h_prev", "Horas previstas de trabalho", "Fluxo"),
    ("h_aus", "Horas de ausência (faltas + atestados)", "Fluxo"),
    ("h_extra", "Horas extras pagas", "Fluxo"),
    ("custo", "Custo de folha com encargos (R$)", "Fluxo"),
    ("vagas_abertas", "Requisições aprovadas e não preenchidas no fim da semana", "Estoque"),
    ("vagas_fech", "Vagas preenchidas na semana", "Fluxo"),
    ("dias_fech", "SOMA dos dias de abertura das vagas fechadas (não a média)", "Fluxo"),
    ("h_treino", "Horas de treinamento realizadas", "Fluxo"),
    ("acid", "Acidentes com afastamento", "Fluxo"),
    ("hc_lid", "Posições de liderança ocupadas", "Estoque"),
    ("hc_lid_fem", "Posições de liderança ocupadas por mulheres", "Estoque"),
    ("enps_resp", "Respostas do pulso de clima", "Fluxo"),
    ("enps_prom", "Promotores no pulso (nota 9 ou 10)", "Fluxo"),
    ("enps_detr", "Detratores no pulso (nota 0 a 6)", "Fluxo"),
]

r = 3
ws.cell(row=2, column=1, value="Componentes que o painel já usa — responda as colunas amarelas").font = F_EX
for nome, desc, tipo in COMPONENTES:
    ws.cell(row=r, column=1, value=nome).font = Font(name="Consolas", size=9, color=TINTA)
    ws.cell(row=r, column=1).fill = P_PRONTO
    ws.cell(row=r, column=2, value=desc).font = F_TXT
    ws.cell(row=r, column=2).fill = P_PRONTO
    ws.cell(row=r, column=3, value=tipo).font = F_TXT
    ws.cell(row=r, column=3).fill = P_PRONTO
    ws.cell(row=r, column=3).alignment = CENTRO
    for j in (4, 5, 6, 7, 8):
        c = ws.cell(row=r, column=j)
        c.font = F_PREENCHER
        c.fill = P_PREENCHER
    for j in range(1, 9):
        ws.cell(row=r, column=j).border = BORDA
        ws.cell(row=r, column=j).alignment = ALTO
    r += 1

ws.cell(row=r + 1, column=1, value="Componentes adicionais que seus indicadores exigem").font = F_EX
extra_ini = r + 2
for rr in range(extra_ini, extra_ini + 12):
    for j in range(1, 9):
        c = ws.cell(row=rr, column=j)
        c.font = F_PREENCHER
        c.fill = P_PREENCHER
        c.border = BORDA
        c.alignment = ALTO

lista(ws, ["Sim", "Não", "Parcial"], "D", 3, extra_ini + 11, "Você tem?")
lista(ws, ["Diária", "Semanal", "Mensal", "Trimestral"], "F", 3, extra_ini + 11, "Granularidade")
lista(ws, ["Matrícula", "Centro de custo", "Área", "Diretoria", "Empresa", "Só o grupo"],
      "G", 3, extra_ini + 11, "Quebra")
lista(ws, ["Fluxo", "Estoque"], "C", extra_ini, extra_ini + 11, "Tipo")

# ================================================================== #
# Aba 3: Hierarquia
# ================================================================== #
ws = wb.create_sheet("3. Hierarquia")
cabecalho(ws, [
    ("#", 4), ("Empresa", 26), ("Diretoria", 26), ("Área", 30),
    ("Headcount aproximado", 18), ("Observações", 40),
])
ws.cell(row=2, column=2, value="⬇ exemplo — uma linha por área, repetindo empresa e diretoria").font = F_EX
for j, v in enumerate(["Indústria Sul", "Operações", "Produção", 612, "Maior área do grupo"], start=2):
    c = ws.cell(row=3, column=j, value=v)
    c.font = F_EX
    c.fill = P_EX
    c.border = BORDA
ws.cell(row=3, column=1, value=1).font = F_EX
for r in range(4, 4 + 60):
    ws.cell(row=r, column=1, value=r - 2).font = F_TXT
    ws.cell(row=r, column=1).alignment = CENTRO
    for j in range(2, 7):
        c = ws.cell(row=r, column=j)
        c.font = F_PREENCHER
        c.fill = P_PREENCHER
        c.border = BORDA
        c.alignment = ALTO

# ================================================================== #
# Aba 4: Tela a tela
# ================================================================== #
ws = wb.create_sheet("4. Tela a tela")
cabecalho(ws, [
    ("#", 4),
    ("Slide / aba de hoje", 30),
    ("O que mostra hoje", 44),
    ("Vira o quê no painel", 30),
    ("Some da apresentação?", 16),
    ("Observações", 40),
])

RASCUNHO = [
    ("Capa e agenda", "Título, semana de referência, sumário", "Cabeçalho", "Sim"),
    ("Resumo executivo", "Bullets do que aconteceu na semana", "Leitura do período", "Sim"),
    ("Headcount — foto", "Quadro atual, admissões e desligamentos", "Cartão de KPI", "Sim"),
    ("Headcount — acumulado mês", "Mesmo indicador, outra janela", "Faixa semana/mês/ano", "Sim"),
    ("Headcount — acumulado ano", "Mesmo indicador, outra janela", "Faixa semana/mês/ano", "Sim"),
    ("Headcount por diretoria", "Mesma métrica quebrada por área", "Filtro", "Sim"),
    ("Turnover — mês", "Rotatividade do mês contra meta", "Cartão de KPI", "Sim"),
    ("Turnover — série 12 meses", "Gráfico de linha do histórico", "Gráfico da análise", "Sim"),
    ("Turnover — comparativo ano anterior", "Duas séries no mesmo gráfico", "Gráfico da análise", "Sim"),
    ("Absenteísmo", "Indicador contra meta", "Cartão de KPI", "Sim"),
    ("Horas extras", "Indicador contra meta", "Cartão de KPI", "Sim"),
    ("Recrutamento — vagas e prazo", "Vagas abertas e tempo de preenchimento", "Cartão de KPI", "Sim"),
    ("Treinamento", "Horas realizadas contra plano", "Cartão de KPI", "Sim"),
    ("Segurança", "Acidentes e taxa de frequência", "Cartão de KPI", "Sim"),
    ("Clima / e-NPS", "Resultado do pulso mais recente", "Cartão de KPI", "Sim"),
    ("Diversidade", "Indicadores de representatividade", "Cartão de KPI", "Sim"),
    ("Comentários e planos de ação", "Texto por indicador", "Motivos e ações", "Sim"),
    ("Anexo — tabelas detalhadas", "Tabela por área", "Tabela da análise", "Sim"),
]
ws.cell(row=2, column=2, value="⬇ rascunho do formato usual — corrija para a sua apresentação real").font = F_EX
r = 3
for i, (slide, mostra, vira, some) in enumerate(RASCUNHO, start=1):
    ws.cell(row=r, column=1, value=i).font = F_TXT
    ws.cell(row=r, column=1).alignment = CENTRO
    for j, v in enumerate([slide, mostra, vira, some], start=2):
        c = ws.cell(row=r, column=j)
        c.value = v
        c.font = F_PREENCHER
        c.fill = P_PREENCHER
        c.alignment = ALTO
    ws.cell(row=r, column=6).fill = P_PREENCHER
    for j in range(1, 7):
        ws.cell(row=r, column=j).border = BORDA
    r += 1
for rr in range(r, r + 22):
    ws.cell(row=rr, column=1, value=rr - 2).font = F_TXT
    ws.cell(row=rr, column=1).alignment = CENTRO
    for j in range(2, 7):
        c = ws.cell(row=rr, column=j)
        c.font = F_PREENCHER
        c.fill = P_PREENCHER
        c.border = BORDA
        c.alignment = ALTO

lista(ws, [
    "Cabeçalho",
    "Leitura do período",
    "Cartão de KPI",
    "Faixa semana/mês/ano",
    "Gráfico da análise",
    "Tabela da análise",
    "Motivos e ações",
    "Filtro",
    "Contador do farol",
    "Sai do escopo",
    "Ainda não sei",
], "D", 3, r + 21, "Vira o quê")
lista(ws, ["Sim", "Não", "Vira anexo"], "E", 3, r + 21, "Some da apresentação?")

# ------------------------------------------------------------------ #
destino = RAIZ / "inventario-indicadores.xlsx"
wb.save(destino)
print(f"{destino.relative_to(RAIZ)} gerado — {len(wb.sheetnames)} abas: {', '.join(wb.sheetnames)}")
